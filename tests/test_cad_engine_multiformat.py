"""Tests for multi-application CAD export adapters."""
import ast
import base64
import json
from io import BytesIO

from openpyxl import load_workbook

from gridpulse.cad_engine import CadAsset, CadRoute, EngineeringSpec, build_spec, review_spec
from gridpulse.cad_engine_multiformat import (
    export_3ds_max_script,
    export_engineering_workbook,
    export_multi_application_bundle,
    export_revit_dynamo_python,
)


def test_multi_application_bundle_uses_one_validated_spec():
    spec, _ = build_spec(
        "industrial power plant with piping and electrical equipment",
        80, 55,
        ["Mechanical", "Piping", "Electrical", "Architecture"],
        "power",
        False,
    )
    bundle = export_multi_application_bundle(spec)
    assert set(bundle) == {
        "autocad_model.dxf",
        "3dsmax_generator.ms",
        "revit_dynamo_generator.py",
        "engineering_spec.json",
        "gridpulse_engineering_handoff.xlsx.base64",
    }
    assert bundle["autocad_model.dxf"].startswith("0\nSECTION")
    assert "GridPulse 3ds Max generator" in bundle["3dsmax_generator.ms"]
    assert "DirectShape.CreateElement" in bundle["revit_dynamo_generator.py"]
    assert json.loads(bundle["engineering_spec.json"])["assets"]
    embedded_xlsx = base64.b64decode(bundle["gridpulse_engineering_handoff.xlsx.base64"])
    assert embedded_xlsx[:2] == b"PK"
    excel = load_workbook(BytesIO(embedded_xlsx), read_only=True, data_only=True)
    assert "Equipment" in excel.sheetnames and "Design Checks" in excel.sheetnames\n    assert "Calculations" in excel.sheetnames and "Rule Checks" in excel.sheetnames
    excel.close()


def test_revit_script_is_valid_python_syntax_and_discloses_limits():
    spec, _ = build_spec("substation", 80, 55, ["Electrical", "Architecture"], "substation", False)
    script = export_revit_dynamo_python(spec)
    ast.parse(script)
    assert "not a native family library" in script
    assert "not auto-converted to MEP systems" in script


def test_3ds_max_script_has_assets_and_routes():
    spec, _ = build_spec("process plant", 80, 55, ["Mechanical", "Piping"], "process", False)
    script = export_3ds_max_script(spec)
    assert "gpObj.name =" in script
    assert "gpRoute.name =" in script
    assert "units.systemType = #meters" in script
    assert "conceptual geometry only" in script.lower()



def test_geometric_review_detects_route_crossing_equipment():
    spec = EngineeringSpec(
        project_name="clash-test",
        site_width=20,
        site_length=20,
        assets=[
            CadAsset(
                tag="P-101", name="Pump", discipline="Mechanical", shape="pump",
                x=5, y=5, z=0, length=4, width=4, height=3,
            )
        ],
        routes=[
            CadRoute(
                tag="P-1001", name="Schematic pipe", discipline="Piping", kind="pipe",
                points=[(1, 6, 1), (10, 6, 1)], width=0.2,
            )
        ],
    )
    report = review_spec(spec)
    clashes = [item for item in report["geometric_findings"]
               if item["code"] == "ROUTE_ASSET_ENVELOPE_CLASH"]
    assert clashes
    assert clashes[0]["items"] == ["P-1001", "P-101"]
    assert report["engineering_approval"] is False


def test_excel_workbook_contains_traceable_multi_sheet_handoff():
    spec, _ = build_spec(
        "industrial power plant with piping and electrical equipment",
        80, 55, ["Mechanical", "Piping", "Electrical", "Architecture"],
        "power", False,
    )
    report = review_spec(spec)
    payload = export_engineering_workbook(spec, report)
    assert payload[:2] == b"PK"
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "Project Summary", "Equipment", "Routes", "Route Points", "Design Checks", "Design Basis", "Calculations", "Rule Checks"
    ]
    assert workbook["Equipment"].max_row == len(spec.assets) + 1
    assert workbook["Routes"].max_row == len(spec.routes) + 1
    assert workbook["Route Points"].max_row == sum(len(r.points) for r in spec.routes) + 1
    assert "NO_ENGINEERING_CALCULATIONS" in {
        row[1] for row in workbook["Design Checks"].iter_rows(min_row=2, values_only=True)
    }
    workbook.close()
